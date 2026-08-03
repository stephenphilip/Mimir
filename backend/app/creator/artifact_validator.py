"""Artifact Validator — mandatory gate before registration."""

from __future__ import annotations

import os
from pathlib import Path
from typing import Optional

from .mime import guess_mime
from .types import ValidationResult
from .diagnostics import get_execution_diagnostics


class ArtifactValidator:
    """Validate artifacts on disk before they may be registered."""

    THUMB_DIR_NAME = ".thumbnails"

    def __init__(self, artifacts_dir: Path) -> None:
        self.artifacts_dir = Path(artifacts_dir)
        self.thumb_dir = self.artifacts_dir / self.THUMB_DIR_NAME
        self.thumb_dir.mkdir(parents=True, exist_ok=True)
        self._diag = get_execution_diagnostics()

    def validate(self, file_path: str, artifact_type: str) -> ValidationResult:
        path = Path(file_path)
        normalized = artifact_type.lower().replace(".", "")

        if normalized == "md":
            normalized = "markdown"

        if not path.exists():
            return self._fail("filesystem", f"File does not exist: {path}")

        if not path.is_file():
            return self._fail("filesystem", f"Path is not a file: {path}")

        size = path.stat().st_size
        if size == 0:
            return self._fail("validation", f"File is empty: {path.name}")

        if normalized == "pdf":
            return self._validate_pdf(path, size)
        if normalized in {"png", "jpg", "jpeg", "webp", "gif", "image"}:
            return self._validate_image(path, size)
        if normalized == "docx":
            return self._validate_docx(path, size)
        if normalized in {"xlsx", "csv"}:
            return self._validate_spreadsheet(path, size, normalized)
        if normalized in {"markdown", "txt"}:
            return self._validate_text(path, size, normalized)

        # Generic fallback: exists, non-empty, MIME guess
        mime = guess_mime(path.name)
        self._diag.log("validation", f"Generic validation passed for {path.name}", metadata={"mime": mime})
        return ValidationResult(valid=True, mime_type=mime, size=size)

    def _fail(self, category: str, message: str) -> ValidationResult:
        self._diag.log(category, message, level="error")
        return ValidationResult(valid=False, errors=[message])

    def _validate_pdf(self, path: Path, size: int) -> ValidationResult:
        try:
            with open(path, "rb") as fh:
                header = fh.read(5)
            if not header.startswith(b"%PDF"):
                return self._fail("validation", f"Invalid PDF header: {path.name}")
            # Attempt minimal read via pypdf if available, else header check suffices
            try:
                from pypdf import PdfReader  # type: ignore

                reader = PdfReader(str(path))
                if len(reader.pages) == 0:
                    return self._fail("validation", f"PDF has no pages: {path.name}")
            except ImportError:
                pass
            except Exception as exc:
                return self._fail("validation", f"PDF not readable: {path.name} ({exc})")

            mime = "application/pdf"
            self._diag.log("validation", f"PDF validated: {path.name}", metadata={"size": size})
            return ValidationResult(valid=True, mime_type=mime, size=size)
        except OSError as exc:
            return self._fail("filesystem", f"Cannot read PDF: {path.name} ({exc})")

    def _validate_image(self, path: Path, size: int) -> ValidationResult:
        try:
            from PIL import Image

            with Image.open(path) as img:
                img.verify()
            with Image.open(path) as img:
                img.load()
                thumb_path = self._make_thumbnail(path, img)
            ext_mime = {
                ".png": "image/png",
                ".jpg": "image/jpeg",
                ".jpeg": "image/jpeg",
                ".webp": "image/webp",
                ".gif": "image/gif",
            }
            mime = ext_mime.get(path.suffix.lower(), guess_mime(path.name))
            self._diag.log(
                "validation",
                f"Image validated: {path.name}",
                metadata={"size": size, "thumbnail": thumb_path},
            )
            return ValidationResult(valid=True, mime_type=mime, size=size, thumbnail_path=thumb_path)
        except ImportError:
            return self._fail("validation", "Pillow required for image validation (pip install Pillow)")
        except Exception as exc:
            return self._fail("validation", f"Image decode failed: {path.name} ({exc})")

    def _make_thumbnail(self, path: Path, img) -> Optional[str]:
        try:
            from PIL import Image

            thumb_name = f"{path.stem}_thumb.jpg"
            thumb_path = self.thumb_dir / thumb_name
            copy = img.copy()
            copy.thumbnail((256, 256))
            if copy.mode in ("RGBA", "P"):
                copy = copy.convert("RGB")
            copy.save(thumb_path, "JPEG", quality=85)
            return str(thumb_path)
        except Exception:
            return None

    def _validate_docx(self, path: Path, size: int) -> ValidationResult:
        try:
            from docx import Document

            doc = Document(str(path))
            _ = len(doc.paragraphs)
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            self._diag.log("validation", f"DOCX validated: {path.name}")
            return ValidationResult(valid=True, mime_type=mime, size=size)
        except ImportError:
            return self._fail("validation", "python-docx required for DOCX validation")
        except Exception as exc:
            return self._fail("validation", f"DOCX open failed: {path.name} ({exc})")

    def _validate_spreadsheet(self, path: Path, size: int, kind: str) -> ValidationResult:
        if kind == "csv":
            try:
                text = path.read_text(encoding="utf-8")
                if not text.strip():
                    return self._fail("validation", f"CSV is empty: {path.name}")
                self._diag.log("validation", f"CSV validated: {path.name}")
                return ValidationResult(valid=True, mime_type="text/csv", size=size)
            except UnicodeDecodeError as exc:
                return self._fail("validation", f"CSV not UTF-8: {path.name} ({exc})")

        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True)
            wb.close()
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            self._diag.log("validation", f"XLSX validated: {path.name}")
            return ValidationResult(valid=True, mime_type=mime, size=size)
        except ImportError:
            return self._fail("validation", "openpyxl required for XLSX validation")
        except Exception as exc:
            return self._fail("validation", f"Workbook not readable: {path.name} ({exc})")

    def _validate_text(self, path: Path, size: int, kind: str) -> ValidationResult:
        try:
            path.read_text(encoding="utf-8")
            mime = "text/markdown" if kind == "markdown" else "text/plain"
            self._diag.log("validation", f"{kind.upper()} validated: {path.name}")
            return ValidationResult(valid=True, mime_type=mime, size=size)
        except UnicodeDecodeError as exc:
            return self._fail("validation", f"{kind} not valid UTF-8: {path.name} ({exc})")
